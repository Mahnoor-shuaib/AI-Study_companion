import os
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    elif ext == ".pdf":
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text

    elif ext == ".docx":
        doc = Document(file_path)
        return "\n\n".join(para.text for para in doc.paragraphs if para.text.strip())

    elif ext == ".pptx":
        prs = Presentation(file_path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            text += run.text + " "
                    text += "\n"
        return text

    else:
        raise ValueError(f"Unsupported file type: {ext}")